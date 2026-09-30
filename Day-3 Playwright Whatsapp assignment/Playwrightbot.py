import json
import time
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright


# ---------------------------------------------------------
# Configuration
# ---------------------------------------------------------

INPUT_FILE = "contacts.xlsx"

print (f"Input file: {INPUT_FILE}")

SCREENSHOT_DIR = Path("screenshots")
REPORT_DIR = Path("reports")

JSON_REPORT = REPORT_DIR / "whatsapp_report.json"
EXCEL_REPORT = REPORT_DIR / "whatsapp_report.xlsx"

WHATSAPP_URL = "https://web.whatsapp.com"


# Create directories
SCREENSHOT_DIR.mkdir(exist_ok=True)
REPORT_DIR.mkdir(exist_ok=True)


# ---------------------------------------------------------
# Read contacts from Excel
# ---------------------------------------------------------

def read_contacts():

    workbook = load_workbook(INPUT_FILE)

    sheet = workbook.active

    contacts = []

    headers = {}

    for cell in sheet[1]:
        if cell.value:
            headers[str(cell.value).strip().lower()] = cell.column

    required_columns = ["name", "phone", "message"]

    for column in required_columns:
        if column not in headers:
            raise Exception(
                f"Missing required column: {column}"
            )

    for row in range(2, sheet.max_row + 1):

        name = sheet.cell(
            row=row,
            column=headers["name"]
        ).value

        phone = sheet.cell(
            row=row,
            column=headers["phone"]
        ).value

        message = sheet.cell(
            row=row,
            column=headers["message"]
        ).value

        if not name or not phone:
            continue

        contacts.append({
            "name": str(name).strip(),
            "phone": str(phone).strip(),
            "message": str(message or "").strip()
        })

    return contacts


# ---------------------------------------------------------
# Replace template variables
# ---------------------------------------------------------

def personalize_message(template, name):

    return template.replace(
        "{name}",
        name
    )


# ---------------------------------------------------------
# Open a WhatsApp chat
# ---------------------------------------------------------

def open_chat(page, phone):

    print(f"Opening chat for {phone}")

    # WhatsApp Web direct chat URL
    chat_url = (
        "https://web.whatsapp.com/send?phone="
        + phone.replace("+", "")
    )

    page.goto(
        chat_url,
        wait_until="domcontentloaded",
        timeout=60000
    )

    page.wait_for_timeout(4000)


# ---------------------------------------------------------
# Send message
# ---------------------------------------------------------

def send_message(page, message):

    print(f"Sending: {message}")

    # WhatsApp message input
    message_box = page.locator(
        'div[contenteditable="true"]'
    ).last

    message_box.wait_for(
        state="visible",
        timeout=30000
    )

    message_box.click()

    message_box.fill(message)

    page.keyboard.press("Enter")

    # Give WhatsApp time to send
    page.wait_for_timeout(2500)

    return True


# ---------------------------------------------------------
# Take screenshot
# ---------------------------------------------------------

def take_screenshot(page, name):

    safe_name = "".join(
        c if c.isalnum() or c in (" ", "_", "-")
        else "_"
        for c in name
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        SCREENSHOT_DIR
        / f"{safe_name}_{timestamp}.png"
    )

    page.screenshot(
        path=str(filename),
        full_page=False
    )

    return str(filename)


# ---------------------------------------------------------
# Extract last 3 messages
# ---------------------------------------------------------

def extract_last_messages(page):

    print("Extracting last 3 messages...")

    page.wait_for_timeout(1500)

    # WhatsApp message containers.
    #
    # These selectors can change when WhatsApp
    # changes its web application.
    message_elements = page.locator(
        'div.message-in, div.message-out'
    )

    count = message_elements.count()

    messages = []

    start = max(0, count - 3)

    for i in range(start, count):

        element = message_elements.nth(i)

        try:
            text = element.inner_text().strip()

            if text:
                messages.append(text)

        except Exception:
            continue

    return messages


# ---------------------------------------------------------
# Save JSON report
# ---------------------------------------------------------

def save_json_report(results):

    with open(
        JSON_REPORT,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            results,
            file,
            indent=4,
            ensure_ascii=False
        )


# ---------------------------------------------------------
# Save Excel report
# ---------------------------------------------------------

def save_excel_report(results):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "WhatsApp Report"

    headers = [
        "Name",
        "Phone",
        "Message Sent",
        "Status",
        "Screenshot",
        "Last Message 1",
        "Last Message 2",
        "Last Message 3",
        "Timestamp"
    ]

    sheet.append(headers)

    for result in results:

        messages = result.get(
            "last_messages",
            []
        )

        sheet.append([
            result.get("name"),
            result.get("phone"),
            result.get("message"),
            result.get("status"),
            result.get("screenshot"),
            messages[0] if len(messages) > 0 else "",
            messages[1] if len(messages) > 1 else "",
            messages[2] if len(messages) > 2 else "",
            result.get("timestamp")
        ])

    workbook.save(EXCEL_REPORT)


# ---------------------------------------------------------
# Main
# ---------------------------------------------------------

def main():

    contacts = read_contacts()

    print(
        f"Loaded {len(contacts)} contacts."
    )

    results = []

    with sync_playwright() as p:

        # Persistent browser profile.
        #
        # This is important because after the first
        # QR-code login, the WhatsApp session can
        # normally be reused.
        context = p.chromium.launch_persistent_context(
            user_data_dir="whatsapp_profile",
            headless=False,
            viewport={
                "width": 1400,
                "height": 900
            }
        )

        page = context.pages[0]

        print("Opening WhatsApp Web...")

        page.goto(
            WHATSAPP_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        print()
        print("=" * 60)
        print("FIRST RUN:")
        print("Scan the WhatsApp QR code manually.")
        print("After login, press ENTER here.")
        print("=" * 60)

        input(
            "\nPress ENTER after WhatsApp Web is logged in..."
        )

        # -------------------------------------------------
        # Process contacts
        # -------------------------------------------------

        for contact in contacts:

            name = contact["name"]
            phone = contact["phone"]
            template = contact["message"]

            timestamp = datetime.now().isoformat()

            personalized_message = personalize_message(
                template,
                name
            )

            result = {
                "name": name,
                "phone": phone,
                "message": personalized_message,
                "status": "Failed",
                "screenshot": "",
                "last_messages": [],
                "timestamp": timestamp
            }

            try:

                print()
                print("=" * 60)
                print(f"Processing: {name}")
                print(f"Phone: {phone}")
                print("=" * 60)

                # Open contact
                open_chat(
                    page,
                    phone
                )

                # Send message
                send_message(
                    page,
                    personalized_message
                )

                # Screenshot
                screenshot = take_screenshot(
                    page,
                    name
                )

                result["screenshot"] = screenshot

                # Extract last 3 messages
                messages = extract_last_messages(
                    page
                )

                result["last_messages"] = messages

                result["status"] = "Success"

                print("Message sent successfully.")
                print("Last messages:")

                for msg in messages:
                    print(
                        f"  {msg}"
                    )

            except Exception as error:

                print(
                    f"ERROR processing {name}:"
                )

                print(error)

                result["status"] = (
                    f"Failed: {str(error)}"
                )

            results.append(result)

            # Small delay between contacts
            time.sleep(2)

        # -------------------------------------------------
        # Save reports
        # -------------------------------------------------

        save_json_report(results)

        save_excel_report(results)

        print()
        print("=" * 60)
        print("PROCESS COMPLETED")
        print("=" * 60)

        print(
            f"JSON report: {JSON_REPORT}"
        )

        print(
            f"Excel report: {EXCEL_REPORT}"
        )

        print(
            f"Screenshots: {SCREENSHOT_DIR}"
        )

        input(
            "\nPress ENTER to close the browser..."
        )

        context.close()


if __name__ == "__main__":
    main()